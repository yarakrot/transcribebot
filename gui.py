import os
import sys
import glob
import threading
import traceback
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# Запуск через pythonw.exe даёт sys.stdout/stderr = None, из-за чего падает tqdm
# (использует его для прогресс-бара внутри whisper.transcribe). Подставляем заглушку.
class _NullWriter:
    def write(self, *_args, **_kwargs):
        pass

    def flush(self):
        pass


if sys.stdout is None:
    sys.stdout = _NullWriter()
if sys.stderr is None:
    sys.stderr = _NullWriter()

# ffmpeg ставится через winget, но PATH в процессе, запущенном из ярлыка,
# может быть не обновлён (Explorer не перезапускался). Добавляем путь напрямую.
def _ensure_ffmpeg_on_path():
    if os.system("where ffmpeg >nul 2>nul") == 0:
        return
    pattern = os.path.expandvars(
        r"%LOCALAPPDATA%\Microsoft\WinGet\Packages\Gyan.FFmpeg_*\ffmpeg-*-full_build\bin"
    )
    matches = glob.glob(pattern)
    if matches:
        os.environ["PATH"] = matches[0] + os.pathsep + os.environ.get("PATH", "")


_ensure_ffmpeg_on_path()

MODELS = ["tiny", "base", "small", "medium", "large-v3"]
LANGUAGES = {
    "Автоопределение": None,
    "Русский": "ru",
    "English": "en",
}

FILETYPES = [
    ("Аудио/Видео", "*.mp3 *.wav *.m4a *.flac *.ogg *.mp4 *.mkv *.mov *.avi *.webm"),
    ("Все файлы", "*.*"),
]


class WhisperGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Whisper — транскрибация записей")
        self.geometry("720x560")
        self.minsize(640, 480)

        self.files = []
        self.model_cache = {}
        self.worker_thread = None
        self.cancel_requested = False

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        # File selection
        frame_files = ttk.LabelFrame(self, text="Файлы для транскрибации")
        frame_files.pack(fill="x", **pad)

        self.files_listbox = tk.Listbox(frame_files, height=6, selectmode=tk.EXTENDED)
        self.files_listbox.pack(fill="x", padx=8, pady=(8, 4), expand=True)

        btn_row = ttk.Frame(frame_files)
        btn_row.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(btn_row, text="Добавить файлы...", command=self.add_files).pack(side="left")
        ttk.Button(btn_row, text="Убрать выбранные", command=self.remove_selected).pack(side="left", padx=6)
        ttk.Button(btn_row, text="Очистить", command=self.clear_files).pack(side="left")

        # Options
        frame_opts = ttk.LabelFrame(self, text="Настройки")
        frame_opts.pack(fill="x", **pad)

        row1 = ttk.Frame(frame_opts)
        row1.pack(fill="x", padx=8, pady=6)
        ttk.Label(row1, text="Модель:").pack(side="left")
        self.model_var = tk.StringVar(value="medium")
        ttk.Combobox(row1, textvariable=self.model_var, values=MODELS, width=12, state="readonly").pack(side="left", padx=(6, 20))

        ttk.Label(row1, text="Язык:").pack(side="left")
        self.lang_var = tk.StringVar(value="Автоопределение")
        ttk.Combobox(row1, textvariable=self.lang_var, values=list(LANGUAGES.keys()), width=16, state="readonly").pack(side="left", padx=6)

        row2 = ttk.Frame(frame_opts)
        row2.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(row2, text="Форматы вывода:").pack(side="left")
        self.fmt_txt = tk.BooleanVar(value=True)
        self.fmt_docx = tk.BooleanVar(value=True)
        self.fmt_srt = tk.BooleanVar(value=False)
        ttk.Checkbutton(row2, text="TXT", variable=self.fmt_txt).pack(side="left", padx=(10, 4))
        ttk.Checkbutton(row2, text="DOCX", variable=self.fmt_docx).pack(side="left", padx=4)
        ttk.Checkbutton(row2, text="SRT (субтитры)", variable=self.fmt_srt).pack(side="left", padx=4)

        row3 = ttk.Frame(frame_opts)
        row3.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(row3, text="Папка сохранения:").pack(side="left")
        self.outdir_var = tk.StringVar(value="")
        ttk.Entry(row3, textvariable=self.outdir_var).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row3, text="Обзор...", command=self.pick_outdir).pack(side="left")
        ttk.Label(frame_opts, text="Если не указана — файлы сохранятся рядом с исходной записью.",
                  foreground="#888").pack(anchor="w", padx=8, pady=(0, 6))

        # Run controls
        frame_run = ttk.Frame(self)
        frame_run.pack(fill="x", **pad)
        self.start_btn = ttk.Button(frame_run, text="Начать транскрибацию", command=self.start)
        self.start_btn.pack(side="left")
        self.cancel_btn = ttk.Button(frame_run, text="Отмена", command=self.cancel, state="disabled")
        self.cancel_btn.pack(side="left", padx=6)

        self.progress = ttk.Progressbar(frame_run, mode="determinate")
        self.progress.pack(side="left", fill="x", expand=True, padx=10)

        # Log
        frame_log = ttk.LabelFrame(self, text="Журнал")
        frame_log.pack(fill="both", expand=True, **pad)
        self.log_text = tk.Text(frame_log, height=12, wrap="word")
        self.log_text.pack(fill="both", expand=True, padx=8, pady=(8, 0))
        self.log_text.bind("<Key>", self._log_block_edit)

        log_btn_row = ttk.Frame(frame_log)
        log_btn_row.pack(fill="x", padx=8, pady=8)
        ttk.Button(log_btn_row, text="Скопировать журнал", command=self.copy_log).pack(side="left")

    # --- file list handlers ---
    def add_files(self):
        paths = filedialog.askopenfilenames(title="Выберите записи", filetypes=FILETYPES)
        for p in paths:
            if p not in self.files:
                self.files.append(p)
                self.files_listbox.insert(tk.END, p)

    def remove_selected(self):
        selected = list(self.files_listbox.curselection())
        for idx in reversed(selected):
            self.files_listbox.delete(idx)
            del self.files[idx]

    def clear_files(self):
        self.files_listbox.delete(0, tk.END)
        self.files.clear()

    def pick_outdir(self):
        d = filedialog.askdirectory(title="Папка сохранения")
        if d:
            self.outdir_var.set(d)

    # --- logging ---
    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)

    def log_threadsafe(self, msg):
        self.after(0, self.log, msg)

    def _log_block_edit(self, event):
        # Allow copy/select-all/navigation, block everything that would type into the log.
        allowed_ctrl = {"c", "a", "insert"}
        if event.state & 0x4 and event.keysym.lower() in allowed_ctrl:
            return None
        if event.keysym in ("Left", "Right", "Up", "Down", "Home", "End", "Prior", "Next", "Shift_L", "Shift_R"):
            return None
        return "break"

    def copy_log(self):
        self.clipboard_clear()
        self.clipboard_append(self.log_text.get("1.0", tk.END))

    # --- run ---
    def start(self):
        if not self.files:
            messagebox.showwarning("Нет файлов", "Добавьте хотя бы один файл.")
            return
        if not (self.fmt_txt.get() or self.fmt_docx.get() or self.fmt_srt.get()):
            messagebox.showwarning("Нет формата", "Выберите хотя бы один формат вывода.")
            return

        self.cancel_requested = False
        self.start_btn.configure(state="disabled")
        self.cancel_btn.configure(state="normal")
        self.progress.configure(maximum=len(self.files), value=0)

        self.worker_thread = threading.Thread(target=self._run_worker, daemon=True)
        self.worker_thread.start()

    def cancel(self):
        self.cancel_requested = True
        self.log("Отмена запрошена — завершится после текущего файла...")

    def _run_worker(self):
        try:
            import whisper
        except Exception:
            self.log_threadsafe("Ошибка импорта whisper:\n" + traceback.format_exc())
            self.after(0, self._on_finish)
            return

        model_name = self.model_var.get()
        lang_code = LANGUAGES.get(self.lang_var.get())
        outdir = self.outdir_var.get().strip() or None

        try:
            if model_name not in self.model_cache:
                self.log_threadsafe(f"Загружаю модель '{model_name}'...")
                self.model_cache[model_name] = whisper.load_model(model_name)
            model = self.model_cache[model_name]
        except Exception:
            self.log_threadsafe("Ошибка загрузки модели:\n" + traceback.format_exc())
            self.after(0, self._on_finish)
            return

        for i, path in enumerate(self.files, start=1):
            if self.cancel_requested:
                break
            self.log_threadsafe(f"\n[{i}/{len(self.files)}] Транскрибирую: {path}")
            try:
                result = model.transcribe(path, language=lang_code, verbose=False)
                self._save_outputs(path, result, outdir)
                self.log_threadsafe(f"Готово: {path}")
            except Exception:
                self.log_threadsafe(f"Ошибка при обработке {path}:\n" + traceback.format_exc())
            self.after(0, lambda v=i: self.progress.configure(value=v))

        self.log_threadsafe("\nВсе задачи завершены." if not self.cancel_requested else "\nОстановлено пользователем.")
        self.after(0, self._on_finish)

    def _on_finish(self):
        self.start_btn.configure(state="normal")
        self.cancel_btn.configure(state="disabled")

    def _save_outputs(self, src_path, result, outdir):
        base = os.path.splitext(os.path.basename(src_path))[0]
        target_dir = outdir or os.path.dirname(src_path) or "."
        os.makedirs(target_dir, exist_ok=True)
        base_path = os.path.join(target_dir, base)
        text = result["text"].strip()

        if self.fmt_txt.get():
            with open(base_path + ".txt", "w", encoding="utf-8") as f:
                f.write(text)
            self.log_threadsafe(f"  -> {base_path}.txt")

        if self.fmt_docx.get():
            from docx import Document
            doc = Document()
            doc.add_heading(base, level=1)
            for para in text.split("\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            doc.save(base_path + ".docx")
            self.log_threadsafe(f"  -> {base_path}.docx")

        if self.fmt_srt.get():
            with open(base_path + ".srt", "w", encoding="utf-8") as f:
                for idx, seg in enumerate(result["segments"], start=1):
                    start = format_timestamp(seg["start"])
                    end = format_timestamp(seg["end"])
                    f.write(f"{idx}\n{start} --> {end}\n{seg['text'].strip()}\n\n")
            self.log_threadsafe(f"  -> {base_path}.srt")


def format_timestamp(seconds: float) -> str:
    ms = int((seconds - int(seconds)) * 1000)
    s = int(seconds)
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    return f"{h:02}:{m:02}:{s:02},{ms:03}"


if __name__ == "__main__":
    app = WhisperGUI()
    app.mainloop()
